"""
exporter.py
-----------
WHAT THIS FILE DOES:
  Takes the finished Markdown string from the LLM and converts it to
  Markdown file, DOCX, or PDF.

TRIGGER CHAIN:
  main.py: export_prd()
      └── to_markdown(content, path)     ← trivial: write string to file
      └── to_docx(content, path)         ← parse Markdown → add to Word doc
      └── to_pdf(content, path)          ← Markdown → PDF via ReportLab
"""

import re
from pathlib import Path


# ── EXPORT 1: Markdown ────────────────────────────────────────────────────────

def to_markdown(content: str, output_path: str) -> str:
    """Writes the Markdown string directly to a .md file."""
    Path(output_path).write_text(content, encoding="utf-8")
    return output_path


# ── EXPORT 2: DOCX ────────────────────────────────────────────────────────────

def to_docx(content: str, output_path: str) -> str:
    """Converts Markdown string → Word .docx file."""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        raise ImportError("Run: pip install python-docx")

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for level, size in [("Heading 1", 20), ("Heading 2", 16), ("Heading 3", 13)]:
        h_style = doc.styles[level]
        h_style.font.size = Pt(size)
        h_style.font.bold = True

    lines = content.split("\n")
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("# ") and not stripped.startswith("## "):
            doc.add_heading(stripped[2:].strip(), level=1)

        elif stripped.startswith("## ") and not stripped.startswith("### "):
            doc.add_heading(stripped[3:].strip(), level=2)

        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)

        elif stripped.startswith(("- [ ]", "- [x]")):
            checked = stripped.startswith("- [x]")
            text = ("✓ " if checked else "☐ ") + stripped[5:].strip()
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_bold(p, text)

        elif stripped.startswith(("- ", "* ")):
            text = stripped[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_bold(p, text)

        elif re.match(r"^\d+\.\s", stripped):
            text = re.sub(r"^\d+\.\s", "", stripped)
            p = doc.add_paragraph(style="List Number")
            _add_runs_with_bold(p, text)

        elif stripped == "---":
            doc.add_paragraph()

        else:
            p = doc.add_paragraph()
            _add_runs_with_bold(p, stripped)

        i += 1

    doc.save(output_path)
    return output_path


def _add_runs_with_bold(paragraph, text: str):
    """Splits text on **...** markers and adds Word runs with bold formatting."""
    parts = re.split(r"\*\*(.*?)\*\*", text)
    for j, part in enumerate(parts):
        if part:
            run = paragraph.add_run(part)
            run.bold = (j % 2 == 1)


# ── EXPORT 3: PDF ─────────────────────────────────────────────────────────────

def to_pdf(content: str, output_path: str) -> str:
    """Converts Markdown string → PDF using ReportLab (works on Windows)."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
    except ImportError:
        raise ImportError("Run: pip install reportlab")

    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )
    styles = getSampleStyleSheet()
    story = []

    for line in content.split("\n"):
        s = line.strip()
        if not s:
            story.append(Spacer(1, 6))
        elif s.startswith("# ") and not s.startswith("## "):
            story.append(Paragraph(s[2:], styles["Title"]))
        elif s.startswith("## ") and not s.startswith("### "):
            story.append(Spacer(1, 10))
            story.append(Paragraph(s[3:], styles["Heading2"]))
        elif s.startswith("### "):
            story.append(Paragraph(s[4:], styles["Heading3"]))
        elif s.startswith(("- ", "* ")):
            story.append(Paragraph(s[2:], styles["Normal"]))
        else:
            story.append(Paragraph(s, styles["Normal"]))

    doc.build(story)
    return output_path